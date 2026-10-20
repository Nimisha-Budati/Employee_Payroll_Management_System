from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.core.cache import cache
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import AttendanceRecord, AuditLog, LeaveRequest, UserProfile
from workforce.models import OvertimeRecord

from .tests import JUNE_2025_PRESENT, check_in, make_employee, make_structure


class ExcelExportTests(TestCase):
    def setUp(self):
        cache.clear()
        self.admin = User.objects.create_user('admin', 'admin@example.com', 'AdminPass123!')
        UserProfile.objects.create(user=self.admin, role='ADMIN')
        self.asha = make_employee('asha', 'EMP001')
        self.ravi = make_employee('ravi', 'EMP002', department='Sales')
        # A malicious value typed into an employee field must never become a live formula.
        self.ravi.designation = '=HYPERLINK("http://evil.example","click")'
        self.ravi.save()
        make_structure(self.asha.user)
        make_structure(self.ravi.user, basic=Decimal('20000'), hra=Decimal('8000'), special_allowance=Decimal('2000'))
        for day in JUNE_2025_PRESENT:
            check_in(self.asha.user, day)
        record = AttendanceRecord.objects.get(employee=self.asha.user, date=date(2025, 6, 2))
        record.late_minutes, record.shift_name = 25, 'General'
        record.save()
        LeaveRequest.objects.create(employee=self.asha.user, leave_type='Casual Leave', start_date=date(2025, 6, 16),
                                    end_date=date(2025, 6, 17), reason='Family', status='APPROVED')
        LeaveRequest.objects.create(employee=self.ravi.user, leave_type='Sick Leave', start_date=date(2025, 6, 18),
                                    end_date=date(2025, 6, 18), reason='Fever', status='PENDING')
        OvertimeRecord.objects.create(employee=self.asha.user, date=date(2025, 6, 5), hours=Decimal('2'), rate=Decimal('300'),
                                      reason='Release', status='APPROVED', source='MANUAL')
        self.client = APIClient()
        self.client.force_authenticate(self.admin)
        self.client.post('/api/payroll/run/', {'year': 2025, 'month': 6}, format='json')

    def download(self, kind, **params):
        response = self.client.get(f'/api/reports/excel/{kind}/', params)
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('.xlsx', response['Content-Disposition'])
        return load_workbook(BytesIO(response.content))

    @staticmethod
    def column(ws, header):
        headers = [cell.value for cell in ws[5]]
        index = headers.index(header)
        return [row[index] for row in ws.iter_rows(min_row=6, values_only=True) if row[0] not in (None, 'Total')]

    def test_employee_report_with_filters_and_formula_safety(self):
        wb = self.download('employees')
        ws = wb['Employees']
        self.assertEqual(self.column(ws, 'Employee ID'), ['EMP001', 'EMP002'])
        designation_cell = next(cell for cell in ws['G'] if cell.value and str(cell.value).startswith('=HYPERLINK'))
        self.assertEqual(designation_cell.data_type, 's')   # stored as text, not a formula
        self.assertIn('By department', wb.sheetnames)

        sales_only = self.download('employees', department='Sales')['Employees']
        self.assertEqual(self.column(sales_only, 'Employee ID'), ['EMP002'])
        self.assertIn('Department: Sales', sales_only['A2'].value)

    def test_attendance_report_includes_late_minutes_and_summary(self):
        wb = self.download('attendance', **{'from': '2025-06-01', 'to': '2025-06-30'})
        records = wb['Attendance']
        self.assertEqual(len(self.column(records, 'Date')), 10)
        self.assertEqual(sum(v or 0 for v in self.column(records, 'Late minutes')), 25)
        summary = wb['Summary']
        self.assertEqual(self.column(summary, 'Late days'), [1])
        self.assertEqual(self.column(summary, 'Days present'), [10])
        late_only = self.download('attendance', **{'from': '2025-06-01', 'to': '2025-06-30', 'late': 'true'})['Attendance']
        self.assertEqual(len(self.column(late_only, 'Date')), 1)
        bad = self.client.get('/api/reports/excel/attendance/', {'from': '2025-06-30', 'to': '2025-06-01'})
        self.assertEqual(bad.status_code, 400)

    def test_leave_report_filters_by_status_and_employee(self):
        wb = self.download('leave', **{'from': '2025-06-01', 'to': '2025-06-30'})
        self.assertEqual(sorted(self.column(wb['Leave requests'], 'Leave type')), ['Casual Leave', 'Sick Leave'])
        approved = self.download('leave', status='APPROVED')['Leave requests']
        self.assertEqual(self.column(approved, 'Name'), ['Asha'])
        ravi_only = self.download('leave', employee=self.ravi.pk)['Leave requests']
        self.assertEqual(self.column(ravi_only, 'Status'), ['Pending'])

    def test_payroll_report_matches_payslips(self):
        wb = self.download('payroll', year=2025, month=6)
        ws = wb['Payroll']
        self.assertEqual(self.column(ws, 'Employee ID'), ['EMP001', 'EMP002'])
        self.assertEqual(self.column(ws, 'Overtime pay')[0], 600)   # approved overtime flows into payroll
        last = ws.max_row
        self.assertEqual(ws.cell(row=last, column=1).value, 'Total')
        self.assertTrue(str(ws.cell(row=last, column=26).value).startswith('=SUM('))
        self.assertEqual(self.column(self.download('payroll', year=2025, month=6, department='Sales')['Payroll'], 'Employee ID'), ['EMP002'])
        self.assertEqual(self.download('overtime', **{'from': '2025-06-01', 'to': '2025-06-30'})['Overtime'].max_row >= 6, True)

    def test_exports_are_admin_only_and_audited(self):
        self.download('payroll', year=2025)
        self.assertTrue(AuditLog.objects.filter(module='Reports', action='EXPORT').exists())
        employee_client = APIClient()
        employee_client.force_authenticate(self.asha.user)
        self.assertEqual(employee_client.get('/api/reports/excel/payroll/').status_code, 403)
        self.assertEqual(self.client.get('/api/reports/excel/unknown/').status_code, 404)
