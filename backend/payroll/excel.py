"""Excel (.xlsx) report exports built with OpenPyXL.

Every report is generated from live database rows with the filters the administrator chose.
Text cells are always stored as text, so a value such as ``=HYPERLINK(...)`` typed into an
employee field can never run as a spreadsheet formula.
"""

import calendar
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from accounts.audit import record_audit
from accounts.models import AttendanceRecord, Employee, LeaveRequest
from accounts.permissions import IsAdminUserRole
from accounts.views import attendance_summary_rows
from workforce.models import OvertimeRecord
from workforce.services import shift_for

from .models import Payslip

GREEN = '047857'
SOFT_GREEN = 'E4F5EC'
MONEY_FORMAT = '"₹"#,##0.00'
THIN = Side(style='thin', color='D1E7DD')

# Column kinds: text, int, number, money, date, datetime, time
EXCEL_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class ReportError(Exception):
    pass


def _local(value):
    if isinstance(value, datetime) and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    return value


def _write_cell(ws, row, column, value, kind):
    if isinstance(value, Decimal):
        value = float(value)
    value = _local(value)
    if kind in ('time',) and isinstance(value, datetime):
        value = value.time()
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        cell.data_type = 's'   # never let user text become a formula
    if kind == 'money':
        cell.number_format = MONEY_FORMAT
    elif kind == 'number':
        cell.number_format = '0.00'
    elif kind == 'date':
        cell.number_format = 'DD-MMM-YYYY'
    elif kind == 'datetime':
        cell.number_format = 'DD-MMM-YYYY HH:MM'
    elif kind == 'time':
        cell.number_format = 'HH:MM'
    return cell


def add_sheet(wb, title, report_title, filters, columns, rows, totals=None):
    """columns: [(header, width, kind)]; rows: iterables matching columns; totals: {column_index: 'sum'|value}."""
    ws = wb.create_sheet(title=title[:31])
    ws.sheet_view.showGridLines = False
    width = len(columns)
    ws.cell(row=1, column=1, value=f'{settings.COMPANY_NAME} – {report_title}').font = Font(bold=True, size=14, color=GREEN)
    ws.cell(row=2, column=1, value=filters).font = Font(italic=True, color='555555')
    ws.cell(row=2, column=1).data_type = 's'
    generated = timezone.localtime().strftime('%d %b %Y %H:%M')
    ws.cell(row=3, column=1, value=f'Generated {generated}').font = Font(color='777777', size=9)

    header_row = 5
    for index, (header, col_width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=index, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.border = Border(bottom=THIN)
        ws.column_dimensions[get_column_letter(index)].width = col_width
    ws.row_dimensions[header_row].height = 30

    current = header_row
    for current, row in enumerate(rows, start=header_row + 1):
        for index, value in enumerate(row, start=1):
            cell = _write_cell(ws, current, index, value, columns[index - 1][2])
            cell.border = Border(bottom=THIN)
            if (current - header_row) % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F6FBF8')
    last_data_row = current
    count = last_data_row - header_row

    if count == 0:
        ws.cell(row=header_row + 1, column=1, value='No records match these filters.').font = Font(italic=True, color='777777')
    elif totals:
        total_row = last_data_row + 1
        ws.cell(row=total_row, column=1, value='Total').font = Font(bold=True)
        for column_index, rule in totals.items():
            letter = get_column_letter(column_index)
            value = f'=SUM({letter}{header_row + 1}:{letter}{last_data_row})' if rule == 'sum' else rule
            cell = ws.cell(row=total_row, column=column_index, value=value)
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor=SOFT_GREEN)
            if columns[column_index - 1][2] == 'money':
                cell.number_format = MONEY_FORMAT
            elif columns[column_index - 1][2] == 'number':
                cell.number_format = '0.00'
        ws.cell(row=total_row, column=1).fill = PatternFill('solid', fgColor=SOFT_GREEN)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if count:
        ws.auto_filter.ref = f'A{header_row}:{get_column_letter(width)}{last_data_row}'
    ws.print_title_rows = f'{header_row}:{header_row}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    return count


# --- filter helpers ------------------------------------------------------------

def _date_param(params, name):
    raw = params.get(name, '') or ''
    if not raw:
        return None
    parsed = parse_date(raw)
    if parsed is None:
        raise ReportError(f'Invalid {name} date. Use YYYY-MM-DD.')
    return parsed


def _range(params, default_start=None, default_end=None):
    start, end = _date_param(params, 'from') or default_start, _date_param(params, 'to') or default_end
    if start and end and end < start:
        raise ReportError('The "to" date cannot be before the "from" date.')
    if start and end and (end - start).days > 366 * 3:
        raise ReportError('Choose a date range of three years or less.')
    return start, end


def _employee_user(params):
    raw = params.get('employee')
    if not raw:
        return None
    employee = Employee.objects.select_related('user').filter(pk=raw).first() if str(raw).isdigit() else None
    if employee is None:
        raise ReportError('Employee not found.')
    return employee.user


def _describe(parts):
    text = ' · '.join(f'{label}: {value}' for label, value in parts if value)
    return f'Filters – {text}' if text else 'Filters – none (all records)'


def _name(user):
    return user.get_full_name() or user.username


def _code(user):
    employee = getattr(user, 'employee', None)
    return employee.employee_id if employee is not None else ''


def _dept(user):
    employee = getattr(user, 'employee', None)
    return employee.department if employee is not None else ''


# --- reports -------------------------------------------------------------------

def employee_report(params, wb, company_id=None):
    employees = Employee.objects.filter(company_id=company_id).select_related('user', 'user__salary_structure').order_by('employee_id')
    if params.get('status') in ('ACTIVE', 'INACTIVE'):
        employees = employees.filter(status=params['status'])
    if params.get('department'):
        employees = employees.filter(department__iexact=params['department'])
    search = (params.get('search') or '').strip()
    if search:
        employees = employees.filter(
            Q(employee_id__icontains=search) | Q(user__username__icontains=search) | Q(user__first_name__icontains=search)
            | Q(user__last_name__icontains=search) | Q(user__email__icontains=search) | Q(designation__icontains=search)
        )
    start, end = _range(params)
    if start:
        employees = employees.filter(joining_date__gte=start)
    if end:
        employees = employees.filter(joining_date__lte=end)

    today = timezone.localdate()
    rows = []
    for employee in employees:
        user = employee.user
        structure = getattr(user, 'salary_structure', None)
        shift = shift_for(user, today)
        rows.append([
            employee.employee_id, _name(user), user.username, user.email, employee.phone, employee.department,
            employee.designation, employee.joining_date, employee.get_status_display(),
            shift.name if shift else '', structure.gross_monthly if structure else None,
            structure.bonus if structure else None,
        ])
    columns = [
        ('Employee ID', 13, 'text'), ('Name', 24, 'text'), ('Username', 16, 'text'), ('Email', 28, 'text'),
        ('Phone', 15, 'text'), ('Department', 18, 'text'), ('Designation', 20, 'text'), ('Joining date', 13, 'date'),
        ('Status', 10, 'text'), ('Shift today', 16, 'text'), ('Monthly gross', 15, 'money'), ('Monthly bonus', 14, 'money'),
    ]
    filters = _describe([('Status', params.get('status')), ('Department', params.get('department')), ('Search', search),
                         ('Joined from', start), ('Joined to', end)])
    count = add_sheet(wb, 'Employees', 'Employee report', filters, columns, rows, totals={11: 'sum', 12: 'sum'})

    departments = {}
    for employee in employees:
        item = departments.setdefault(employee.department, [employee.department, 0, 0, 0])
        item[1] += 1
        item[2 if employee.status == 'ACTIVE' else 3] += 1
    add_sheet(wb, 'By department', 'Headcount by department', filters,
              [('Department', 24, 'text'), ('Employees', 12, 'int'), ('Active', 10, 'int'), ('Inactive', 10, 'int')],
              sorted(departments.values()), totals={2: 'sum', 3: 'sum', 4: 'sum'})
    return count, filters


def attendance_report(params, wb, company_id=None):
    today = timezone.localdate()
    start, end = _range(params, today.replace(day=1), today)
    records = AttendanceRecord.objects.select_related('employee', 'employee__employee').filter(employee__userprofile__company_id=company_id, date__range=(start, end))
    if params.get('department'):
        records = records.filter(employee__employee__department__iexact=params['department'])
    user = _employee_user(params)
    if user is not None:
        records = records.filter(employee=user)
    if params.get('late') == 'true':
        records = records.filter(late_minutes__gt=0)
    records = records.order_by('date', 'employee__employee__employee_id', 'employee__username')

    rows = []
    for record in records:
        worked = None
        if record.check_in_at and record.check_out_at:
            worked = round((record.check_out_at - record.check_in_at).total_seconds() / 3600, 2)
        rows.append([
            record.date, _code(record.employee), _name(record.employee), _dept(record.employee), record.shift_name,
            record.scheduled_start, record.check_in_at, record.check_out_at, worked, record.late_minutes,
            'Checked out' if record.check_out_at else 'Checked in',
            record.check_in_distance_meters, record.check_out_distance_meters,
        ])
    columns = [
        ('Date', 13, 'date'), ('Employee ID', 13, 'text'), ('Name', 22, 'text'), ('Department', 16, 'text'),
        ('Shift', 14, 'text'), ('Scheduled start', 11, 'time'), ('Check in', 10, 'time'), ('Check out', 10, 'time'),
        ('Hours worked', 10, 'number'), ('Late minutes', 10, 'int'), ('Status', 12, 'text'),
        ('Check-in distance (m)', 12, 'int'), ('Check-out distance (m)', 12, 'int'),
    ]
    filters = _describe([('From', start), ('To', end), ('Department', params.get('department')),
                         ('Employee', _name(user) if user else ''), ('Only late', 'yes' if params.get('late') == 'true' else '')])
    count = add_sheet(wb, 'Attendance', 'Attendance report', filters, columns, rows, totals={9: 'sum', 10: 'sum'})

    summary = attendance_summary_rows(records)
    add_sheet(wb, 'Summary', 'Attendance summary by employee', filters, [
        ('Employee ID', 13, 'text'), ('Name', 22, 'text'), ('Department', 16, 'text'), ('Days present', 10, 'int'),
        ('Late days', 10, 'int'), ('Late minutes', 11, 'int'), ('Average late (min)', 11, 'int'), ('Hours worked', 11, 'number'),
    ], [[item['employee_code'], item['employee_name'], item['department'], item['days_present'], item['late_days'],
         item['late_minutes'], item['average_late_minutes'], round(item['worked_minutes'] / 60, 2)]
        for item in summary['employees']], totals={4: 'sum', 5: 'sum', 6: 'sum', 8: 'sum'})
    return count, filters


def leave_report(params, wb, company_id=None):
    start, end = _range(params)
    leaves = LeaveRequest.objects.select_related('employee', 'employee__employee', 'policy').filter(employee__userprofile__company_id=company_id)
    if start:
        leaves = leaves.filter(end_date__gte=start)
    if end:
        leaves = leaves.filter(start_date__lte=end)
    if params.get('status') in ('PENDING', 'APPROVED', 'REJECTED'):
        leaves = leaves.filter(status=params['status'])
    if params.get('leave_type'):
        leaves = leaves.filter(leave_type=params['leave_type'])
    if params.get('department'):
        leaves = leaves.filter(employee__employee__department__iexact=params['department'])
    user = _employee_user(params)
    if user is not None:
        leaves = leaves.filter(employee=user)
    leaves = leaves.order_by('start_date', 'employee__username')

    rows = [[
        _code(item.employee), _name(item.employee), _dept(item.employee), item.leave_type, item.start_date, item.end_date,
        item.number_of_days, item.get_status_display(), ('Paid' if item.policy.is_paid else 'Unpaid') if item.policy_id else '',
        item.reason, item.rejection_reason or '', item.applied_at,
    ] for item in leaves]
    columns = [
        ('Employee ID', 13, 'text'), ('Name', 22, 'text'), ('Department', 16, 'text'), ('Leave type', 14, 'text'),
        ('From', 12, 'date'), ('To', 12, 'date'), ('Days', 7, 'int'), ('Status', 11, 'text'), ('Paid/unpaid', 10, 'text'),
        ('Reason', 34, 'text'), ('Admin response', 28, 'text'), ('Applied on', 17, 'datetime'),
    ]
    filters = _describe([('From', start), ('To', end), ('Status', params.get('status')), ('Type', params.get('leave_type')),
                         ('Department', params.get('department')), ('Employee', _name(user) if user else '')])
    count = add_sheet(wb, 'Leave requests', 'Leave report', filters, columns, rows, totals={7: 'sum'})

    from accounts.models import LeaveType

    types = list(LeaveType.objects.filter(company_id=company_id).values_list('name', flat=True))
    types += sorted({item.leave_type for item in leaves} - set(types))
    per_employee = {}
    for item in leaves:
        key = item.employee_id
        entry = per_employee.setdefault(key, {'row': [_code(item.employee), _name(item.employee), _dept(item.employee)],
                                              'approved': dict.fromkeys(types, 0), 'pending': 0, 'rejected': 0})
        if item.status == 'APPROVED':
            entry['approved'][item.leave_type] = entry['approved'].get(item.leave_type, 0) + item.number_of_days
        elif item.status == 'PENDING':
            entry['pending'] += item.number_of_days
        else:
            entry['rejected'] += item.number_of_days
    summary_rows = [entry['row'] + [entry['approved'][t] for t in types] + [entry['pending'], entry['rejected']]
                    for entry in per_employee.values()]
    summary_columns = [('Employee ID', 13, 'text'), ('Name', 22, 'text'), ('Department', 16, 'text')] + \
        [(f'{t} (approved days)', 14, 'int') for t in types] + [('Pending days', 11, 'int'), ('Rejected days', 11, 'int')]
    add_sheet(wb, 'Summary', 'Leave days by employee', filters, summary_columns, summary_rows,
              totals={index: 'sum' for index in range(4, len(summary_columns) + 1)})
    return count, filters


PAYROLL_COLUMNS = [
    ('Period', 11, 'text'), ('Employee ID', 13, 'text'), ('Name', 22, 'text'), ('Department', 16, 'text'),
    ('Designation', 18, 'text'), ('Working days', 9, 'int'), ('LOP days', 8, 'number'), ('Paid days', 8, 'number'),
    ('Basic', 13, 'money'), ('HRA', 12, 'money'), ('DA', 12, 'money'), ('Conveyance', 12, 'money'), ('Medical', 12, 'money'),
    ('Special', 12, 'money'), ('Other allowance', 12, 'money'), ('Overtime hours', 9, 'number'), ('Overtime pay', 12, 'money'),
    ('Bonus', 12, 'money'), ('Gross', 14, 'money'), ('PF', 12, 'money'), ('Professional tax', 12, 'money'),
    ('TDS', 12, 'money'), ('Other deductions', 12, 'money'), ('LOP deduction', 12, 'money'),
    ('Total deductions', 14, 'money'), ('Net salary', 14, 'money'), ('Status', 9, 'text'), ('Paid on', 13, 'date'),
]


def payroll_report(params, wb, company_id=None):
    try:
        year = int(params.get('year') or timezone.localdate().year)
        month = int(params['month']) if params.get('month') else None
    except ValueError:
        raise ReportError('Provide a valid year and month.')
    if month is not None and not 1 <= month <= 12:
        raise ReportError('Provide a valid month.')
    payslips = Payslip.objects.filter(employee__userprofile__company_id=company_id, year=year)
    if month:
        payslips = payslips.filter(month=month)
    if params.get('status') in (Payslip.DRAFT, Payslip.PAID):
        payslips = payslips.filter(status=params['status'])
    if params.get('department'):
        payslips = payslips.filter(department__iexact=params['department'])
    user = _employee_user(params)
    if user is not None:
        payslips = payslips.filter(employee=user)
    payslips = payslips.order_by('month', 'employee_code')

    rows = [[
        f'{calendar.month_abbr[s.month]} {s.year}', s.employee_code, s.employee_name, s.department, s.designation,
        s.working_days, s.lop_days, max(Decimal('0'), Decimal(s.working_days) - s.lop_days), s.basic, s.hra, s.da,
        s.conveyance_allowance, s.medical_allowance, s.special_allowance, s.other_allowance, s.overtime_hours,
        s.overtime_pay, s.bonus, s.gross_earnings, s.pf, s.professional_tax, s.tds, s.other_deduction, s.lop_deduction,
        s.total_deductions, s.net_salary, s.get_status_display(), s.paid_at.date() if s.paid_at else None,
    ] for s in payslips]
    period = f'{calendar.month_name[month]} {year}' if month else str(year)
    filters = _describe([('Period', period), ('Status', params.get('status')), ('Department', params.get('department')),
                         ('Employee', _name(user) if user else '')])
    count = add_sheet(wb, 'Payroll', 'Payroll register', filters, PAYROLL_COLUMNS, rows,
                      totals={index: 'sum' for index in range(9, 27)})

    departments = {}
    for s in payslips:
        item = departments.setdefault(s.department or 'Unassigned', [s.department or 'Unassigned', 0, Decimal('0'), Decimal('0'), Decimal('0')])
        item[1] += 1
        item[2] += s.gross_earnings
        item[3] += s.total_deductions
        item[4] += s.net_salary
    add_sheet(wb, 'By department', 'Payroll cost by department', filters,
              [('Department', 22, 'text'), ('Payslips', 10, 'int'), ('Gross', 15, 'money'), ('Deductions', 15, 'money'), ('Net', 15, 'money')],
              sorted(departments.values(), key=lambda item: -item[4]), totals={2: 'sum', 3: 'sum', 4: 'sum', 5: 'sum'})
    return count, filters


def overtime_report(params, wb, company_id=None):
    today = timezone.localdate()
    start, end = _range(params, today.replace(day=1), today)
    records = OvertimeRecord.objects.select_related('employee', 'employee__employee', 'reviewed_by').filter(employee__userprofile__company_id=company_id, date__range=(start, end))
    if params.get('status') in dict(OvertimeRecord.STATUS_CHOICES):
        records = records.filter(status=params['status'])
    if params.get('department'):
        records = records.filter(employee__employee__department__iexact=params['department'])
    user = _employee_user(params)
    if user is not None:
        records = records.filter(employee=user)
    records = records.order_by('date', 'employee__username')
    rows = [[
        r.date, _code(r.employee), _name(r.employee), _dept(r.employee), r.hours, r.rate, r.amount, r.get_status_display(),
        r.get_source_display(), r.reason, r.review_note, r.reviewed_by.username if r.reviewed_by_id else '',
    ] for r in records]
    columns = [
        ('Date', 12, 'date'), ('Employee ID', 13, 'text'), ('Name', 22, 'text'), ('Department', 16, 'text'),
        ('Hours', 8, 'number'), ('Rate / hour', 11, 'money'), ('Amount', 13, 'money'), ('Status', 10, 'text'),
        ('Source', 20, 'text'), ('Reason', 30, 'text'), ('Review note', 24, 'text'), ('Reviewed by', 14, 'text'),
    ]
    filters = _describe([('From', start), ('To', end), ('Status', params.get('status')), ('Department', params.get('department')),
                         ('Employee', _name(user) if user else '')])
    count = add_sheet(wb, 'Overtime', 'Overtime report', filters, columns, rows, totals={5: 'sum', 7: 'sum'})
    return count, filters


def analytics_workbook(params, wb, company_id=None):
    from .analytics import AnalyticsError, build_analytics

    try:
        data = build_analytics(params, company_id)
    except AnalyticsError as error:
        raise ReportError(str(error))
    filters = _describe([('Months', f'{data["from"]} to {data["to"]}'), ('Department', data['department'] or 'All')])
    monthly = [(
        row['label'], row['payslips'], row['payroll']['gross'], row['payroll']['net'], row['payroll']['deductions'],
        row['overtime']['hours'], row['overtime']['amount'], row['attendance']['present'], row['attendance']['late'],
        row['attendance']['late_minutes'], row['leave_days'], row['absenteeism']['lop_days'],
        row['absenteeism']['rate'] if row['absenteeism']['rate'] is not None else '',
        row['headcount']['end'], row['headcount']['joined'], row['headcount']['left'],
    ) for row in data['monthly']]
    count = add_sheet(wb, 'Monthly trends', 'Analytics – monthly trends', filters, [
        ('Month', 12, 'text'), ('Payslips', 10, 'int'), ('Payroll cost (gross)', 18, 'money'), ('Net pay', 16, 'money'),
        ('Deductions', 16, 'money'), ('Overtime hours', 12, 'number'), ('Overtime cost', 15, 'money'),
        ('Check-ins', 11, 'int'), ('Late check-ins', 12, 'int'), ('Late minutes', 12, 'int'), ('Leave days', 11, 'int'),
        ('LOP days', 10, 'number'), ('Absenteeism %', 13, 'number'), ('Headcount (month end)', 14, 'int'),
        ('Joined', 9, 'int'), ('Left', 9, 'int'),
    ], monthly, totals={2: 'sum', 3: 'sum', 4: 'sum', 5: 'sum', 6: 'sum', 7: 'sum', 8: 'sum', 9: 'sum', 10: 'sum',
                        11: 'sum', 12: 'sum', 15: 'sum', 16: 'sum'})
    add_sheet(wb, 'Departments', 'Analytics – departments', filters, [
        ('Department', 22, 'text'), ('Active headcount', 14, 'int'), ('Payslips', 10, 'int'), ('Payroll cost (gross)', 18, 'money'),
        ('Net pay', 16, 'money'), ('Overtime cost', 15, 'money'), ('Late check-ins', 12, 'int'), ('Leave days', 11, 'int'),
    ], [(row['department'], row['headcount'], row['payslips'], row['gross'], row['net'], row['overtime_amount'], row['late'],
         row['leave_days']) for row in data['departments']], totals={2: 'sum', 3: 'sum', 4: 'sum', 5: 'sum', 6: 'sum', 7: 'sum', 8: 'sum'})
    utilisation = data['leave_utilisation']
    add_sheet(wb, 'Leave utilisation', f'Leave utilisation – {utilisation["label"]}', filters, [
        ('Leave type', 20, 'text'), ('Paid', 8, 'text'), ('Entitled days', 13, 'int'), ('Used days', 11, 'int'),
        ('Pending days', 12, 'int'), ('Utilisation %', 13, 'number'),
    ], [(row['leave_type'], 'Yes' if row['is_paid'] else 'No', row['entitled_days'] if row['entitled_days'] is not None else 'No limit',
         row['used_days'], row['pending_days'], row['utilisation'] if row['utilisation'] is not None else '')
        for row in utilisation['types']])
    return count, filters


REPORTS = {
    'analytics': ('Analytics', analytics_workbook),
    'employees': ('Employee', employee_report),
    'attendance': ('Attendance', attendance_report),
    'leave': ('Leave', leave_report),
    'payroll': ('Payroll', payroll_report),
    'overtime': ('Overtime', overtime_report),
}


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsAdminUserRole])
def excel_report(request, kind):
    if kind not in REPORTS:
        return Response({'detail': 'Unknown report type.'}, status=status.HTTP_404_NOT_FOUND)
    label, builder = REPORTS[kind]
    wb = Workbook()
    wb.remove(wb.active)
    try:
        count, filters = builder(request.query_params, wb, request.user.userprofile.company_id)
    except ReportError as error:
        return Response({'detail': str(error)}, status=status.HTTP_400_BAD_REQUEST)
    wb.properties.creator = settings.COMPANY_NAME
    wb.properties.title = f'{label} report'
    buffer = BytesIO()
    wb.save(buffer)
    record_audit(request, 'EXPORT', 'Report', kind, f'Exported {label.lower()} report to Excel ({count} row(s)). {filters}')
    response = HttpResponse(buffer.getvalue(), content_type=EXCEL_CONTENT_TYPE)
    stamp = timezone.localtime().strftime('%Y%m%d-%H%M')
    response['Content-Disposition'] = f'attachment; filename="{kind}_report_{stamp}.xlsx"'
    response['Cache-Control'] = 'private, no-store'
    return response
