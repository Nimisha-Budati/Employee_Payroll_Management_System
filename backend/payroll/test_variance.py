"""Payroll variance review, the review requirement before finalising, analytics and the analytics Excel export."""

from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.contrib.auth.models import User
from django.core.cache import cache
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import AttendanceCorrectionRequest, AttendanceRecord, AuditLog, Employee, LeaveRequest, UserProfile
from workforce.models import OvertimeRecord

from .calculator import working_dates
from .models import Payslip, PayrollReview, SalaryStructure
from .tests import make_employee, make_structure


def check_in(user, day, hour=9, minute=0, late=None):
    start = timezone.make_aware(datetime.combine(day, time(hour, minute)))
    return AttendanceRecord.objects.create(employee=user, date=day, check_in_at=start, check_out_at=start + timedelta(hours=9),
                                           late_minutes=late)


class PayrollCase(TestCase):
    def setUp(self):
        cache.clear()
        self.admin = User.objects.create_user('admin', 'admin@example.com', 'AdminPass123!')
        UserProfile.objects.create(user=self.admin, role='ADMIN')
        self.client = APIClient()
        self.client.force_authenticate(self.admin)
        self.asha = make_employee('asha', 'EMP001', department='Engineering')
        self.ravi = make_employee('ravi', 'EMP002', department='Engineering')
        self.meena = make_employee('meena', 'EMP003', department='Sales')
        for employee in (self.asha, self.ravi, self.meena):
            make_structure(employee.user)

    def run_payroll(self, year, month, mode='NONE'):
        response = self.client.post('/api/payroll/run/', {'year': year, 'month': month, 'lop_mode': mode}, format='json')
        self.assertIn(response.status_code, (200, 201), response.data)
        return response

    def review(self, year, month, **extra):
        return self.client.post('/api/payroll/variance/review/', {'year': year, 'month': month, **extra}, format='json')

    def finalize(self, year, month):
        return self.client.post('/api/payroll/finalize/', {'year': year, 'month': month}, format='json')


class VarianceTests(PayrollCase):
    def setUp(self):
        super().setUp()
        self.run_payroll(2025, 8)
        self.review(2025, 8, acknowledge_flags=True)
        self.finalize(2025, 8)
        # September: Asha gets a raise, Meena loses her salary structure (so she is not paid), Ravi is unchanged.
        SalaryStructure.objects.filter(employee=self.asha.user).update(basic=Decimal('40000'))
        SalaryStructure.objects.filter(employee=self.meena.user).delete()
        self.run_payroll(2025, 9)

    def rows(self, **params):
        data = self.client.get('/api/payroll/variance/', {'year': 2025, 'month': 9, **params}).data
        return data, {row['employee_code']: row for row in data['rows']}

    def test_flags_and_explanations_from_records(self):
        data, rows = self.rows()
        self.assertEqual(data['previous_period'], 'August 2025')
        self.assertEqual(data['counts']['drafts'], 2)
        asha = rows['EMP001']
        self.assertEqual(asha['severity'], 'medium')
        self.assertIn('NET_CHANGE', [f['code'] for f in asha['flags']])
        self.assertIn('STRUCTURE', [f['code'] for f in asha['flags']])
        self.assertTrue(any('Basic ₹30,000 → ₹40,000' in reason for reason in asha['reasons']))
        basic = next(c for c in asha['components'] if c['field'] == 'basic')
        self.assertEqual((Decimal(basic['previous']), Decimal(basic['current']), Decimal(basic['change'])),
                         (Decimal('30000'), Decimal('40000'), Decimal('10000')))
        self.assertEqual(Decimal(basic['percent']), Decimal('33.3'))
        self.assertEqual(rows['EMP002']['flags'], [])
        meena = rows['EMP003']
        self.assertEqual(meena['flags'][0]['code'], 'MISSING')
        self.assertIn('no salary structure', meena['flags'][0]['message'])
        flagged, _ = self.rows(flagged='true')
        self.assertEqual({row['employee_code'] for row in flagged['rows']}, {'EMP001', 'EMP003'})
        # Nothing was changed by looking at the preview.
        self.assertEqual(Payslip.objects.filter(year=2025, month=9, status=Payslip.DRAFT).count(), 2)

    def test_leavers_are_listed_as_information_so_totals_reconcile(self):
        Employee.objects.filter(pk=self.meena.pk).update(status='INACTIVE', deactivated_at=date(2025, 9, 3))
        data, rows = self.rows()
        meena = rows['EMP003']
        self.assertEqual([f['code'] for f in meena['flags']], ['LEFT'])
        self.assertIn('Deactivated on 03 Sep 2025', meena['flags'][0]['message'])
        self.assertFalse(meena['anomaly'])
        self.assertEqual(Decimal(data['totals']['net_salary']['previous']),
                         sum((Decimal(r['previous_net']) for r in data['rows'] if r['previous_net'] is not None), Decimal('0')))

    def test_review_is_required_and_expires_when_drafts_change(self):
        refused = self.finalize(2025, 9)
        self.assertEqual(refused.status_code, 409)
        self.assertEqual(refused.data['code'], 'review_required')
        self.assertEqual(self.review(2025, 9).status_code, 400)            # two flagged employees need acknowledging
        self.assertEqual(self.review(2025, 9, note='Raise approved by HR').status_code, 201)
        self.assertTrue(AuditLog.objects.filter(entity_type='PayrollReview', action='REVIEW', module='Payroll').exists())

        slip = Payslip.objects.get(employee=self.ravi.user, year=2025, month=9)
        Payslip.objects.filter(pk=slip.pk).update(updated_at=timezone.now() + timedelta(seconds=5))   # edited after the review
        self.assertEqual(self.finalize(2025, 9).status_code, 409)
        self.assertFalse(self.client.get('/api/payroll/variance/', {'year': 2025, 'month': 9}).data['review']['is_current'])
        self.assertEqual(self.review(2025, 9, acknowledge_flags=True).status_code, 201)
        PayrollReview.objects.filter(year=2025, month=9).update(reviewed_at=timezone.now() + timedelta(seconds=10))
        done = self.finalize(2025, 9)
        self.assertEqual(done.status_code, 200, done.data)
        self.assertEqual(done.data['finalized'], 2)

    def test_manual_and_stale_drafts_are_flagged(self):
        slip = Payslip.objects.get(employee=self.ravi.user, year=2025, month=9)
        self.client.patch(f'/api/payroll/payslips/{slip.pk}/', {'bonus': '2500'}, format='json')
        _, rows = self.rows()
        self.assertIn('MANUAL', [f['code'] for f in rows['EMP002']['flags']])
        self.assertTrue(any('salary structure bonus is ₹0' in r for r in rows['EMP002']['reasons']))

        OvertimeRecord.objects.create(employee=self.asha.user, date=date(2025, 9, 10), hours=Decimal('4'), rate=Decimal('500'),
                                      status=OvertimeRecord.APPROVED)
        _, rows = self.rows()
        stale = next(f for f in rows['EMP001']['flags'] if f['code'] == 'STALE')
        self.assertIn('approved overtime', stale['message'])

    def test_loss_of_pay_reason_and_permissions(self):
        october = sorted(working_dates(2025, 10))
        for day in october[2:]:
            check_in(self.ravi.user, day)
        self.run_payroll(2025, 10, mode='ATTENDANCE')
        data = self.client.get('/api/payroll/variance/', {'year': 2025, 'month': 10}).data
        ravi = next(row for row in data['rows'] if row['employee_code'] == 'EMP002')
        self.assertTrue(any('2 working day(s) with no check-in, approved leave or WFH' in reason for reason in ravi['reasons']), ravi['reasons'])

        employee = APIClient()
        employee.force_authenticate(self.ravi.user)
        self.assertEqual(employee.get('/api/payroll/variance/', {'year': 2025, 'month': 9}).status_code, 403)
        self.assertEqual(employee.post('/api/payroll/variance/review/', {'year': 2025, 'month': 9}, format='json').status_code, 403)
        self.assertEqual(self.client.get('/api/payroll/variance/', {'year': 2025, 'month': 13}).status_code, 400)


class AnalyticsTests(PayrollCase):
    def setUp(self):
        super().setUp()
        self.run_payroll(2025, 8)
        self.run_payroll(2025, 9)
        Employee.objects.filter(pk=self.meena.pk).update(status='INACTIVE', deactivated_at=date(2025, 9, 15))
        self.new = make_employee('kiran', 'EMP004', joining=date(2025, 9, 1), department='=HYPERLINK("http://x")')
        check_in(self.asha.user, date(2025, 9, 2), late=12)
        check_in(self.ravi.user, date(2025, 9, 2), late=0)
        LeaveRequest.objects.create(employee=self.asha.user, leave_type='Casual Leave', start_date=date(2025, 8, 29),
                                    end_date=date(2025, 9, 2), reason='trip', status='APPROVED')
        OvertimeRecord.objects.create(employee=self.ravi.user, date=date(2025, 9, 5), hours=Decimal('2'), rate=Decimal('400'),
                                      status=OvertimeRecord.APPROVED)

    def test_trends_comparison_and_departments(self):
        data = self.client.get('/api/reports/analytics/', {'from': '2025-08', 'to': '2025-09'}).data
        aug, sep = data['monthly']
        self.assertEqual(aug['payslips'], 3)
        self.assertEqual(Decimal(aug['payroll']['gross']), Decimal('150000'))
        self.assertEqual(aug['leave_days'], 3)           # 29–31 Aug
        self.assertEqual(sep['leave_days'], 2)           # 1–2 Sep
        self.assertEqual((sep['attendance']['present'], sep['attendance']['late']), (2, 1))
        self.assertEqual(Decimal(sep['attendance']['late_rate']), Decimal('50.0'))
        self.assertEqual(Decimal(sep['overtime']['amount']), Decimal('800'))
        self.assertEqual(sep['headcount'], {'end': 3, 'joined': 1, 'left': 1})
        self.assertEqual(aug['headcount']['end'], 3)
        self.assertEqual(data['previous_from'], '2025-06')
        self.assertIsNone(data['comparison']['payroll_cost']['percent'])     # nothing to compare with
        self.assertEqual(Decimal(data['summary']['payroll_cost']), Decimal('300000'))
        engineering = next(d for d in data['departments'] if d['department'] == 'Engineering')
        self.assertEqual(engineering['late'], 1)
        self.assertEqual(Decimal(engineering['overtime_amount']), Decimal('800'))
        self.assertIn('Some months include draft payslips', ' '.join(data['notes']))
        sales = self.client.get('/api/reports/analytics/', {'from': '2025-09', 'to': '2025-09', 'department': 'Sales'}).data
        self.assertEqual(sales['monthly'][0]['payslips'], 1)
        self.assertEqual(sales['monthly'][0]['headcount']['left'], 1)
        utilisation = {row['leave_type']: row for row in data['leave_utilisation']['types']}
        self.assertIn('Casual Leave', utilisation)

    def test_validation_permissions_and_excel(self):
        self.assertEqual(self.client.get('/api/reports/analytics/', {'from': '2025-09', 'to': '2025-08'}).status_code, 400)
        self.assertEqual(self.client.get('/api/reports/analytics/', {'from': '2020-01', 'to': '2025-08'}).status_code, 400)
        self.assertEqual(self.client.get('/api/reports/analytics/', {'from': 'x'}).status_code, 400)
        employee = APIClient()
        employee.force_authenticate(self.asha.user)
        self.assertEqual(employee.get('/api/reports/analytics/').status_code, 403)

        response = self.client.get('/api/reports/excel/analytics/', {'from': '2025-08', 'to': '2025-09'})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Monthly trends', 'Departments', 'Leave utilisation'])
        departments = workbook['Departments']
        cells = [row[0] for row in departments.iter_rows(min_row=6) if row[0].value]
        injected = next(cell for cell in cells if str(cell.value).startswith('=HYPERLINK'))
        self.assertEqual(injected.data_type, 's')      # stored as text, never as a formula
        self.assertTrue(AuditLog.objects.filter(entity_type='Report', entity_id='analytics').exists())
        # CSV export keeps working.
        csv = self.client.get('/api/payroll/export/', {'year': 2025, 'month': 9})
        self.assertEqual(csv.status_code, 200)
        self.assertIn('text/csv', csv['Content-Type'])


@override_settings(GEMINI_API_KEY='test-key', GEMINI_MODEL='gemini-test', AI_MAX_REQUESTS_PER_HOUR=40)
class AssistantNewTopicTests(PayrollCase):
    def test_corrections_onboarding_and_leave_rules_are_grounded_in_own_records(self):
        from assistant.tests import gemini_reply

        day = timezone.localdate() - timedelta(days=2)
        AttendanceCorrectionRequest.objects.create(employee=self.asha.user, date=day, correction_type='MISSED_CHECK_OUT',
                                                   reason='asha gate issue', review_note='Approved by HR', status='APPROVED')
        AttendanceCorrectionRequest.objects.create(employee=self.ravi.user, date=day, correction_type='MISSED_DAY',
                                                   reason='ravi secret reason')
        client = APIClient()
        client.force_authenticate(self.asha.user)
        conversation = client.post('/api/ai/conversations/', {'feature': 'CHAT'}, format='json').data['id']
        with mock.patch('assistant.gemini.requests.post', return_value=gemini_reply()) as post:
            client.post(f'/api/ai/conversations/{conversation}/messages/',
                        {'message': 'What happened to my attendance correction and my onboarding checklist?'}, format='json')
            client.post(f'/api/ai/conversations/{conversation}/messages/', {'message': 'How many days of my leave are left?'}, format='json')
        first = post.call_args_list[0].kwargs['json']['system_instruction']['parts'][0]['text']
        self.assertIn('ATTENDANCE CORRECTION REQUESTS', first)
        self.assertIn('Approved by HR', first)
        self.assertIn('ONBOARDING CHECKLIST: in progress', first)
        self.assertNotIn('ravi secret reason', first)
        self.assertIn('name the record', first)
        second = post.call_args_list[1].kwargs['json']['system_instruction']['parts'][0]['text']
        self.assertIn('LEAVE RULES (set by administrators)', second)
